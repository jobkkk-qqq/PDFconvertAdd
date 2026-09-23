#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
电子发票识别模块
提取PDF发票关键字段（发票类型/号码/日期/购销方/金额/明细），支持文字层和OCR
"""

import os
import re
import tempfile
import fitz  # PyMuPDF

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

# ============================================================
# 发票类型检测
# ============================================================

INVOICE_TYPE_PATTERNS = [
    (r'全面数字化的电子发票|数电发票|数电票', '全面数字化的电子发票（数电票）'),
    (r'电子发票（增值税专用发票）', '电子发票（专用发票）'),
    (r'电子发票（增值税普通发票）', '电子发票（普通发票）'),
    (r'电子发票（专用发票）|增值税电子专用发票', '电子发票（专用发票）'),
    (r'电子发票（普通发票）|增值税电子普通发票', '电子发票（普通发票）'),
    (r'增值税专用发票', '增值税专用发票'),
    (r'增值税普通发票', '增值税普通发票'),
]

# 发票抬头（顶部标题原文）匹配：优先消费主标题形态
TITLE_PATTERN = re.compile(
    r'全面数字化的电子发票（?[^（()）]*）?'
    r'|电子发票[（(][^（()）]*[)）]'
    r'|增值税电子[专普]用发票'
    r'|增值税[专普]用发票'
)


def detect_invoice_type(text):
    """检测发票类型"""
    for pattern, name in INVOICE_TYPE_PATTERNS:
        if re.search(pattern, text):
            return name
    return None


# ============================================================
# 文本提取（文字层优先，图片型走OCR）
# ============================================================

def extract_invoice_text(pdf_path, log_func=None):
    """
    提取PDF文本。
    返回: (text, used_ocr: bool)
    文字层为空或极少 → 用OCR（若可用）
    """
    text = ""
    doc = fitz.open(pdf_path)
    for page in doc:
        text += page.get_text() + "\n"
    doc.close()

    # 判断是否有足够文字层
    meaningful = [l for l in text.split('\n') if l.strip()]
    if len(meaningful) >= 8 and len(text.strip()) > 100:
        return text, False

    # 尝试OCR
    try:
        from converter_gui import get_ocr_engine, ocr_image_with_layout
        if log_func:
            log_func("检测到文字层缺失，使用OCR识别...")
        engine_ready = get_ocr_engine(log_func)
        ocr_text = ""
        doc = fitz.open(pdf_path)
        tmp_dir = tempfile.gettempdir()
        for i, page in enumerate(doc):
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            if pix.n > 3:
                pix = fitz.Pixmap(fitz.csRGB, pix)
            tmp_img = os.path.join(tmp_dir, f"invoice_ocr_{i}.png")
            pix.save(tmp_img)
            items = ocr_image_with_layout(tmp_img, log_func)
            # 按 y 聚类排序
            items.sort(key=lambda it: (int(it['y_center'] / 50), it['x_center']))
            for it in items:
                ocr_text += it['text'] + "\n"
            try:
                os.remove(tmp_img)
            except Exception:
                pass
        doc.close()
        return ocr_text, True
    except ImportError:
        return text, False
    except Exception as e:
        if log_func:
            log_func(f"OCR识别失败: {e}")
        return text, False


# ============================================================
# 字段解析
# ============================================================

def _norm(text):
    """规范化：去空白"""
    return re.sub(r'\s+', '', text)


def _find_block(text, start_kw, end_kw=None):
    """提取两个关键字之间的文本块"""
    idx = text.find(start_kw)
    if idx < 0:
        return ""
    idx += len(start_kw)
    if end_kw:
        end = text.find(end_kw, idx)
        if end < 0:
            end = len(text)
        return text[idx:end]
    return text[idx:]


def _extract_name_taxid(block):
    """从信息块提取 名称 和 纳税人识别号"""
    name, taxid = None, None
    # 名称：名称：xxx 或 名称:xxx
    m = re.search(r'名\s*称\s*[:：]\s*([^\n]{1,60})', block)
    if m:
        name = m.group(1).strip()
    # 税号：纳税人识别号：xxx 或 统一社会信用代码：xxx
    m = re.search(r'(?:纳税人识别号|统一社会信用代码)\s*[:：]\s*([0-9A-Za-z]{15,20})', block)
    if m:
        taxid = m.group(1).strip()
    return name, taxid


def parse_invoice(text):
    """解析发票字段"""
    result = {
        'invoice_type': detect_invoice_type(text),
        'invoice_title': None,     # 发票抬头（顶部标题原文）
        'invoice_code': None,      # 发票代码（旧版）
        'invoice_number': None,   # 发票号码
        'issue_date': None,       # 开票日期
        'buyer_name': None, 'buyer_taxid': None,
        'seller_name': None, 'seller_taxid': None,
        'total_amount': None,     # 价税合计小写
        'total_amount_upper': None,  # 价税合计大写
        'check_code': None,       # 校验码
        'remark': None,           # 备注
        'drawer': None,           # 开票人
        'items': [],              # 项目明细 [{name, spec, unit, qty, price, amount, tax_rate, tax}]
    }

    # 发票抬头（顶部标题原文）
    m = TITLE_PATTERN.search(text)
    if m:
        result['invoice_title'] = re.sub(r'\s+', '', m.group(0))

    # 发票号码（数电票20位 / 旧版8位）
    m = re.search(r'发票号码\s*[:：]?\s*([0-9]{20}|[0-9]{8})', text)
    if m:
        result['invoice_number'] = m.group(1)
    else:
        # 数电票无标签，直接找20位数字（排除税号）
        m = re.search(r'(?<!\d)([0-9]{20})(?!\d)', text)
        if m:
            result['invoice_number'] = m.group(1)

    # 发票代码（旧版12位）
    m = re.search(r'发票代码\s*[:：]?\s*([0-9]{10,12})', text)
    if m:
        result['invoice_code'] = m.group(1)

    # 开票日期
    m = re.search(r'开票日期\s*[:：]?\s*(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日', text)
    if not m:
        m = re.search(r'开票日期\s*[:：]?\s*(\d{4})[-/](\d{1,2})[-/](\d{1,2})', text)
    if m:
        result['issue_date'] = f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"

    # 购买方/销售方
    buyer_block = _find_block(text, '购买方', '销售方')
    if buyer_block:
        result['buyer_name'], result['buyer_taxid'] = _extract_name_taxid(buyer_block)
    else:
        # 旧版：购 买 方 / 名称：xxx
        m = re.search(r'购\s*买\s*方[\s\S]{0,10}?名\s*称\s*[:：]\s*([^\n]{1,60})', text)
        if m:
            result['buyer_name'] = m.group(1).strip()
        m = re.search(r'(?:购\s*买\s*方|购买方信息)[\s\S]{0,80}?纳税人识别号\s*[:：]\s*([0-9A-Za-z]{15,20})', text)
        if m:
            result['buyer_taxid'] = m.group(1)

    seller_block = _find_block(text, '销售方')
    if seller_block:
        result['seller_name'], result['seller_taxid'] = _extract_name_taxid(seller_block)
    else:
        m = re.search(r'销\s*售\s*方[\s\S]{0,10}?名\s*称\s*[:：]\s*([^\n]{1,60})', text)
        if m:
            result['seller_name'] = m.group(1).strip()
        m = re.search(r'(?:销\s*售\s*方|销售方信息)[\s\S]{0,80}?纳税人识别号\s*[:：]\s*([0-9A-Za-z]{15,20})', text)
        if m:
            result['seller_taxid'] = m.group(1)

    # 价税合计（小写）
    m = re.search(r'价税合计\s*[（(]?\s*大写\s*[)）]?[\s\S]{0,60}?[¥￥]\s*([0-9,]+\.?\d{0,2})', text)
    if not m:
        m = re.search(r'[¥￥]\s*([0-9,]+\.\d{2})', text)
    if m:
        result['total_amount'] = m.group(1)

    # 价税合计（大写）
    m = re.search(r'价税合计\s*[（(]\s*大写\s*[)）]\s*([^\n¥￥]{2,30})', text)
    if not m:
        m = re.search(r'([零壹贰叁肆伍陆柒捌玖拾佰仟万亿元角分整]{4,30})', text)
    if m:
        result['total_amount_upper'] = m.group(1).strip()

    # 校验码
    m = re.search(r'校验码\s*[:：]?\s*([0-9]+)', text)
    if m:
        result['check_code'] = m.group(1)

    # 备注
    m = re.search(r'备\s*注\s*[:：]?\s*([^\n]{1,100})', text)
    if m:
        result['remark'] = m.group(1).strip()

    # 开票人
    m = re.search(r'开\s*票\s*人\s*[:：]?\s*([\u4e00-\u9fffA-Za-z·]{1,20})', text)
    if m:
        result['drawer'] = m.group(1).strip()

    # 项目明细：定位表头（项目名称/货物或应税劳务、服务名称）
    result['items'] = parse_items(text)

    return result


def extract_layout_lines(pdf_path):
    """逐页提取每个文本 span 及其坐标。返回 [{'x0','y0','x1','y1','yc','text'}]"""
    lines = []
    doc = fitz.open(pdf_path)
    try:
        for page in doc:
            td = page.get_text("dict")
            for b in td["blocks"]:
                if b.get("type", 0) != 0:      # 跳过图片块
                    continue
                for l in b["lines"]:
                    for sp in l["spans"]:
                        t = re.sub(r'\s+', '', sp.get("text", ""))
                        if not t:
                            continue
                        sx0, sy0, sx1, sy1 = sp["bbox"]
                        lines.append({"x0": sx0, "y0": sy0, "x1": sx1, "y1": sy1,
                                      "yc": (sy0 + sy1) / 2.0, "text": t})
    finally:
        doc.close()
    return lines


def _same_line(a, b, tol=7.0):
    return abs(a["yc"] - b["yc"]) <= tol


def _right_value(lines, anchor, tol=7.0):
    """同行中位于 anchor 右侧、离 anchor 右缘最近的文本 span"""
    cands = [it for it in lines if it is not anchor and _same_line(anchor, it, tol)
             and it["x0"] >= anchor["x0"]]
    if not cands:
        return None
    return min(cands, key=lambda c: abs(c["x0"] - anchor["x1"]))


def _column_x(lines, chars, tol_x=10.0, min_hits=3):
    """检测由 charset 组成的标题（可能为竖排单字）所在的横坐标列表。
    只有该坐标命中标题首字（如“购/销”）且命中多次才算栏，避免左右栏共享字干扰。"""
    from collections import defaultdict
    full = "".join(chars)
    single = set(chars)
    first = chars[0]
    hits = defaultdict(int)
    has_first = set()
    for it in lines:
        t = it["text"]
        if t == full:
            hits[int(it["x0"])] += 100
            has_first.add(int(it["x0"]))
        elif len(t) == 1 and t in single:
            if t == first:
                has_first.add(int(it["x0"]))
            hits[int(it["x0"])] += 1
    return sorted(x for x, n in hits.items() if n >= min_hits and x in has_first)


def _page_width(lines):
    return max([it["x1"] for it in lines], default=500.0)


def _assign_col(lbl_x, buy_x, sell_x, mid):
    """按横坐标把标签归到购买方/销售方栏（左右布局）"""
    if buy_x and sell_x:
        return "buy" if abs(lbl_x - buy_x) <= abs(lbl_x - sell_x) else "sell"
    if buy_x and not sell_x:
        return "buy"
    if sell_x and not buy_x:
        return "sell"
    return "buy" if lbl_x < mid else "sell"


def parse_invoice_layout(lines):
    """基于坐标版面解析发票字段（买卖方按栏、明细按表头列）。
    无法识别基本栏位时返回 None，交由文本流解析回退。"""
    if not lines:
        return None
    from collections import defaultdict
    whole = "".join(it["text"] for it in lines)
    mid = _page_width(lines) / 2.0

    result = {"invoice_type": detect_invoice_type(whole), "invoice_code": None,
              "invoice_title": None, "invoice_number": None, "issue_date": None,
              "buyer_name": None, "buyer_taxid": None,
              "seller_name": None, "seller_taxid": None,
              "total_amount": None, "total_amount_upper": None,
              "check_code": None, "remark": None, "drawer": None, "items": []}

    buy_x = _column_x(lines, "购买方信息")
    sell_x = _column_x(lines, "销售方信息")
    if not buy_x and not sell_x:
        return None

    # ---- 名称 / 税号（按栏归属）----
    name_anchors, tax_anchors = [], []
    for it in lines:
        t = it["text"]
        if re.fullmatch(r'名称\s*[:：]?', t):
            name_anchors.append(it)
        elif '统一社会信用代码' in t or '纳税人识别号' in t:
            tax_anchors.append(it)

    def _gather_name(anchor):
        first = _right_value(lines, anchor)
        if not first:
            return None
        nm, nx, by = first["text"], first["x0"], first["yc"]
        # 补自动折行的续行（紧邻下一行、同栏缩进文字、非标签）
        cands = [it for it in lines
                 if 0.5 < it["y0"] - by < 40
                 and abs(it["x0"] - nx) <= 24
                 and not re.search(r'识别号|纳税人|名称', it["text"])
                 and re.search(r'[\u4e00-\u9fffA-Za-z*（]', it["text"])]
        if cands:
            cands.sort(key=lambda c: c["y0"])
            nm += cands[0]["text"]
        return nm

    def _gather_tax(anchor):
        v = _right_value(lines, anchor, tol=10.0)
        if v:
            m = re.search(r'([0-9A-Z]{15,20})', v["text"])
            if m:
                return m.group(1)
        for it in lines:
            if it is not anchor and _same_line(anchor, it, 10) \
                    and it["x0"] >= anchor["x1"] - 8:
                m = re.search(r'([0-9A-Z]{15,20})', it["text"])
                if m:
                    return m.group(1)
        return None

    bxs = buy_x[0] if buy_x else None
    sxs = sell_x[0] if sell_x else None

    # ---- 发票抬头（顶部标题原文）----
    page_mid_y = max([it["y1"] for it in lines], default=800.0) / 2.0
    title_spans = [it for it in lines
                   if TITLE_PATTERN.search(it["text"]) and it["y0"] < page_mid_y]
    if title_spans:
        title_spans.sort(key=lambda c: c["y0"])
        result["invoice_title"] = re.sub(r'\s+', '', title_spans[0]["text"])

    for a in name_anchors:
        which = _assign_col(a["x0"], bxs, sxs, mid)
        nm = _gather_name(a)
        if nm:
            if which == "buy":
                result["buyer_name"] = nm
            else:
                result["seller_name"] = nm
    for a in tax_anchors:
        which = _assign_col(a["x0"], bxs, sxs, mid)
        tx = _gather_tax(a)
        if tx:
            if which == "buy":
                result["buyer_taxid"] = tx
            else:
                result["seller_taxid"] = tx

    # ---- 发票号码 ----
    for it in lines:
        if re.fullmatch(r'\d{20}', it["text"]):
            result["invoice_number"] = it["text"]
            break
    if not result["invoice_number"]:
        for it in lines:
            if '发票号码' in it["text"]:
                v = _right_value(lines, it)
                if v and re.fullmatch(r'\d{6,20}', v["text"]):
                    result["invoice_number"] = v["text"]
                break

    # ---- 发票代码（旧版10-12位，位于"发票代码："标签右侧）----
    if not result["invoice_code"]:
        for it in lines:
            if '发票代码' in it["text"]:
                v = _right_value(lines, it)
                if v:
                    m = re.search(r'\d{10,12}', v["text"])
                    if m:
                        result["invoice_code"] = m.group(1)
                break

    # ---- 开票日期 ----
    for it in lines:
        if '开票日期' in it["text"]:
            v = _right_value(lines, it)
            if v:
                m = re.search(r'(\d{4})年(\d{1,2})月(\d{1,2})日', v["text"])
                if m:
                    result["issue_date"] = f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
            break
    if not result["issue_date"]:
        for it in lines:
            m = re.search(r'^(\d{4})年(\d{1,2})月(\d{1,2})日$', it["text"])
            if m:
                result["issue_date"] = f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
                break

    # ---- 价税合计（小写）----
    for it in lines:
        if '小写' in it["text"]:
            row = [x for x in lines if _same_line(x, it) and x["x0"] >= it["x1"] - 4]
            nums = [m for x in row for m in re.findall(r'[0-9][0-9,]*\.\d{1,2}', x["text"])]
            if nums:
                result["total_amount"] = nums[0].replace(',', '')
            break
    if not result["total_amount"]:
        for it in lines:
            if '价税合计' in it["text"] and '小写' in it["text"]:
                v = _right_value(lines, it)
                if v:
                    m = re.search(r'[0-9][0-9,]*\.\d{1,2}', v["text"])
                    if m:
                        result["total_amount"] = m.group(1).replace(',', '')
                break

    # ---- 价税合计（大写）----
    for it in lines:
        if '大写' in it["text"] and '价税合计' in it["text"]:
            v = _right_value(lines, it)
            if v and re.search(r'[零壹贰叁肆伍陆柒捌玖拾佰仟万亿元角分整]', v["text"]):
                result["total_amount_upper"] = v["text"]
            break

    # ---- 校验码 ----
    for it in lines:
        m = re.search(r'校验码\s*[:：]?\s*([0-9]+)', it["text"])
        if m:
            result["check_code"] = m.group(1)

    # ---- 备注（左下角“备注”栏右侧第一段文本）----
    for it in lines:
        if it["text"] == "备":
            bx, by = it["x0"], it["y0"]
            cand = [x for x in lines if x["text"] not in ("备", "注")
                    and x["x0"] > bx + 6 and x["y0"] >= by - 3]
            if cand:
                cand.sort(key=lambda c: (c["y0"], c["x0"]))
                result["remark"] = cand[0]["text"]
            break

    # ---- 开票人 ----
    for it in lines:
        if '开票人' in it["text"]:
            v = _right_value(lines, it, tol=10.0)
            if v and re.search(r'[\u4e00-\u9fffA-Za-z·]', v["text"]):
                result["drawer"] = re.sub(r'[:：]', '', v["text"]).strip()
            break

    # ---- 项目明细（按表头列定位）----
    _parse_items_coord(lines, result)

    return result


def _parse_items_coord(lines, result):
    """按表头列位置解析项目明细行"""
    header_y = None
    for it in lines:
        if it["text"] == '项目名称':
            header_y = it["yc"]
            break
    if header_y is None:
        return
    col_map = {}
    name_to_col = {"项目名称": "name", "规格型号": "spec", "单位": "unit", "数量": "qty",
                   "单价": "price", "金额": "amount", "税率": "tax_rate", "税额": "tax"}
    for it in lines:
        if not (abs(it["yc"] - header_y) <= 9):
            continue
        t = it["text"]
        for k, f in name_to_col.items():
            if k in t and f not in col_map:
                col_map[f] = (it["x0"] + it["x1"]) / 2.0
                break
    if not col_map:
        return
    end_y = None
    for it in lines:
        if it["text"] in ("合计", "合计："):
            end_y = it["yc"]
            break
    if end_y is None:
        for it in lines:
            if it["text"] == "合":
                end_y = it["yc"]
                break

    # 按 y 聚合成行、按与列中线最近分配列
    groups = []
    for it in lines:
        if it["text"] in ("合", "计"):
            continue
        if it["yc"] <= header_y + 3 or (end_y and it["yc"] >= end_y - 2):
            continue
        best, bd = None, 1e9
        for f, midx in col_map.items():
            d = abs((it["x0"] + it["x1"]) / 2.0 - midx)
            if d < bd:
                bd, best = d, f
        key = None
        for g in groups:
            if abs(g["yc"] - it["yc"]) <= 7:
                key = g
                break
        if key is None:
            key = {"yc": it["yc"], "cols": {}}
            groups.append(key)
        key["cols"].setdefault(best, []).append(it)

    items = []
    for g in sorted(groups, key=lambda g: g["yc"]):
        cols = g["cols"]
        if 'name' not in cols and 'amount' not in cols and 'price' not in cols:
            continue
        item = {"name": "", "spec": "", "unit": "", "qty": "",
                "price": "", "amount": "", "tax_rate": "", "tax": ""}
        for field in item:
            lst = cols.get(field) or []
            if not lst:
                continue
            lst.sort(key=lambda c: c["x0"])
            val = re.sub(r'[¥￥]', '', ''.join(c["text"] for c in lst))
            item[field] = val
        items.append(item)
        if len(items) >= 50:
            break
    result["items"] = items


def parse_items(text):
    """解析项目明细行"""
    items = []
    # 精确定位表头：项目名称 / 货物或应税劳务、服务名称 / 货物或应税劳务名称
    header_m = None
    for kw in ['项目名称', '货物或应税劳务、服务名称', '货物或应税劳务名称', '服务名称', '规格型号']:
        header_m = re.search(kw, text)
        if header_m:
            break
    if not header_m:
        return items

    start = header_m.end()
    # 明细区域：表头之后到"价税合计"或"合计"
    region = text[start:]
    end_m = re.search(r'价\s*税\s*合\s*计|合\s*计', region)
    if end_m:
        region = region[:end_m.start()]

    # 每行分割
    lines = [l.strip() for l in region.split('\n') if l.strip()]
    # 过滤信息块行（购/销方信息等）
    skip_kw = ['名称', '识别号', '纳税人', '地址', '电话', '开户行', '账号', '购买方', '销售方']

    for line in lines:
        line_clean = re.sub(r'\s+', ' ', line)
        # 跳过信息块行
        if any(k in line_clean for k in skip_kw):
            continue
        # 提取数字列
        nums = re.findall(r'[\d,]+\.?\d*%?', line_clean)
        if len(nums) < 2:
            continue
        item = {'name': '', 'spec': '', 'unit': '', 'qty': '', 'price': '',
                'amount': '', 'tax_rate': '', 'tax': ''}
        # 名称：去掉数字列后的文本（去掉序号）
        text_part = re.sub(r'[\d,]+\.?\d*%?', '|', line_clean)
        parts = [p.strip() for p in text_part.split('|') if p.strip()]
        if parts:
            # 去掉首位序号（如"1."）
            first = parts[0]
            first = re.sub(r'^[0-9]{1,3}[.、]', '', first)
            # 从名称末尾剥离单位（套/个/次/台...）
            for u in ['套', '个', '次', '台', '件', '批', '米', '千克', '公斤', '小时',
                      '天', '月', '年', '箱', '张', '份', '组', '只', '双', '条', '辆', '项']:
                if first.endswith(u) and len(first) > len(u):
                    item['unit'] = u
                    first = first[:-len(u)].strip()
                    break
            item['name'] = first[:40]
            if len(parts) > 1:
                item['spec'] = parts[1][:20]
            if len(parts) > 2 and not item['unit']:
                item['unit'] = parts[2][:10]
        # 数字列按发票标准顺序: 数量 单价 金额 税率 税额
        plain = [n for n in nums if not n.endswith('%')]
        pct = [n for n in nums if n.endswith('%')]
        if pct:
            item['tax_rate'] = pct[0]
        if len(plain) >= 4:
            item['qty'], item['price'], item['amount'], item['tax'] = plain[0], plain[1], plain[2], plain[3]
        elif len(plain) == 3:
            item['qty'], item['price'], item['amount'] = plain[0], plain[1], plain[2]
        elif len(plain) == 2:
            item['price'], item['amount'] = plain[0], plain[1]
        elif len(plain) == 1:
            item['amount'] = plain[0]
        items.append(item)
    return items[:50]  # 最多50行


# ============================================================
# 保存Excel
# ============================================================

def save_invoice_excel(data, output_path, log_func=None):
    """保存识别结果到Excel（Sheet1发票信息 + Sheet2项目明细）"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))
    center = Alignment(horizontal='center', vertical='center')

    # Sheet1: 发票信息
    ws = wb.active
    ws.title = "发票信息"
    fields = [
        ('发票类型', data.get('invoice_type') or '未识别'),
        ('发票抬头', data.get('invoice_title') or '未识别'),
        ('发票代码', data.get('invoice_code') or '未识别'),
        ('发票号码', data.get('invoice_number') or '未识别'),
        ('开票日期', data.get('issue_date') or '未识别'),
        ('购买方名称', data.get('buyer_name') or '未识别'),
        ('购买方纳税人识别号', data.get('buyer_taxid') or '未识别'),
        ('销售方名称', data.get('seller_name') or '未识别'),
        ('销售方纳税人识别号', data.get('seller_taxid') or '未识别'),
        ('价税合计(小写)', data.get('total_amount') or '未识别'),
        ('价税合计(大写)', data.get('total_amount_upper') or '未识别'),
        ('校验码', data.get('check_code') or '未识别'),
        ('备注', data.get('remark') or '未识别'),
        ('开票人', data.get('drawer') or '未识别'),
    ]
    for r, (k, v) in enumerate(fields, 1):
        c1 = ws.cell(row=r, column=1, value=k)
        c2 = ws.cell(row=r, column=2, value=v)
        c1.font = Font(bold=True)
        c1.border = thin_border
        c2.border = thin_border
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 45

    # Sheet2: 项目明细
    ws2 = wb.create_sheet(title="项目明细")
    headers = ['项目名称', '规格型号', '单位', '数量', '单价', '金额', '税率', '税额']
    for c, h in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin_border

    items = data.get('items', [])
    for r, item in enumerate(items, 2):
        vals = [item.get('name', ''), item.get('spec', ''), item.get('unit', ''),
                item.get('qty', ''), item.get('price', ''), item.get('amount', ''),
                item.get('tax_rate', ''), item.get('tax', '')]
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(row=r, column=c, value=v)
            cell.border = thin_border

    for c in range(1, len(headers) + 1):
        ws2.column_dimensions[get_column_letter(c)].width = 18

    wb.save(output_path)
    if log_func:
        log_func(f"✓ 识别结果已保存: {output_path}")
    return output_path


# ============================================================
# 识别窗口
# ============================================================

class InvoiceRecognizerWindow:
    """电子发票识别窗口"""

    def __init__(self, parent_root, checker, log_func, initial_file=None):
        self.parent = parent_root
        self.checker = checker
        self.log = log_func

        self.win = tk.Toplevel(parent_root)
        self.win.title("电子发票识别")
        self.win.geometry("720x600")
        self.win.transient(parent_root)

        self.result_data = None
        self.output_dir = os.path.dirname(initial_file) if initial_file else os.path.expanduser("~")

        self.create_widgets()

        if initial_file and os.path.exists(initial_file):
            self.file_var.set(initial_file)

    def create_widgets(self):
        win = self.win

        # 顶部：文件选择
        top = ttk.Frame(win, padding="8")
        top.pack(fill=tk.X)
        ttk.Label(top, text="发票PDF:").pack(side=tk.LEFT)
        self.file_var = tk.StringVar()
        entry = ttk.Entry(top, textvariable=self.file_var, width=45)
        entry.pack(side=tk.LEFT, padx=5)
        ttk.Button(top, text="浏览...", command=self.choose_file).pack(side=tk.LEFT, padx=3)
        try:
            from tkinterdnd2 import DND_FILES
            entry.drop_target_register(DND_FILES)
            entry.dnd_bind('<<Drop>>', lambda e: self._drop(e))
        except Exception:
            pass
        ttk.Button(top, text="开始识别", command=self.recognize, width=12)\
            .pack(side=tk.LEFT, padx=8)

        # 结果区
        result_frame = ttk.LabelFrame(win, text="识别结果（可选中复制）", padding="8")
        result_frame.pack(fill=tk.BOTH, expand=True, padx=8, pady=5)

        self.result_text = tk.Text(result_frame, height=20, width=80, wrap=tk.WORD)
        self.result_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar = ttk.Scrollbar(result_frame, orient=tk.VERTICAL, command=self.result_text.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.result_text['yscrollcommand'] = scrollbar.set
        self.result_text.config(state='disabled')

        # 底部操作
        bottom = ttk.Frame(win, padding="8")
        bottom.pack(fill=tk.X)

        out_frame = ttk.Frame(bottom)
        out_frame.pack(fill=tk.X)
        ttk.Label(out_frame, text="保存目录:").pack(side=tk.LEFT)
        self.out_var = tk.StringVar(value=self.output_dir)
        out_entry = ttk.Entry(out_frame, textvariable=self.out_var, width=45)
        out_entry.pack(side=tk.LEFT, padx=5)
        ttk.Button(out_frame, text="浏览...", command=self.choose_output).pack(side=tk.LEFT)

        btn_frame = ttk.Frame(bottom)
        btn_frame.pack(fill=tk.X, pady=(8, 0))
        ttk.Button(btn_frame, text="复制全部结果", command=self.copy_result, width=14)\
            .pack(side=tk.LEFT, padx=3)
        self.save_btn = ttk.Button(btn_frame, text="保存为Excel", command=self.save_excel, width=14,
                                   state='disabled')
        self.save_btn.pack(side=tk.LEFT, padx=3)

    def choose_file(self):
        path = filedialog.askopenfilename(
            title="选择发票PDF", filetypes=[("PDF文件", "*.pdf"), ("所有文件", "*.*")])
        if path:
            self.file_var.set(path)
            self.output_dir = os.path.dirname(path)
            self.out_var.set(self.output_dir)

    def _drop(self, event):
        try:
            files = self.win.tk.splitlist(event.data)
            if files:
                self.file_var.set(files[0])
        except Exception:
            pass

    def choose_output(self):
        folder = filedialog.askdirectory(title="选择保存目录")
        if folder:
            self.out_var.set(folder)

    def recognize(self):
        """执行识别"""
        pdf_path = self.file_var.get().strip()
        if not pdf_path:
            messagebox.showwarning("提示", "请选择发票PDF文件")
            return
        if not os.path.exists(pdf_path):
            messagebox.showerror("错误", f"文件不存在: {pdf_path}")
            return

        if not self.checker.is_registered():
            messagebox.showwarning("警告", "请先注册")
            return
        usage = self.checker.get_usage()
        if usage.get('files_remaining', 0) <= 0:
            messagebox.showerror("错误", "配额已用完，请联系开发者")
            return

        self.log(f"开始发票识别: {os.path.basename(pdf_path)}")
        try:
            text, used_ocr = extract_invoice_text(pdf_path, self.log)
            data = None
            # 文字层可用时：优先按版面坐标解析（买卖方按栏、明细按列），更准确
            if not used_ocr:
                try:
                    layout_lines = extract_layout_lines(pdf_path)
                    data = parse_invoice_layout(layout_lines)
                    if data:
                        self.log("  ✓ 已按版面坐标解析（购买方/销售方按栏位、明细按表头列）")
                except Exception as e:
                    self.log(f"坐标解析失败，回退文本解析: {e}")
                    data = None
            if not data:
                data = parse_invoice(text)
            else:
                # 坐标解析缺字段时，用文本流解析补缺
                tdata = parse_invoice(text)
                for k, v in tdata.items():
                    if k != 'items' and not data.get(k) and v:
                        data[k] = v
                if not data.get('items'):
                    data['items'] = tdata.get('items', [])
            self.result_data = data

            # 检查是否发票
            if not data['invoice_type'] and not data['invoice_number']:
                self.log("未识别到发票特征")
                messagebox.showwarning("提示", "未识别到发票特征，请确认文件是电子发票PDF")

            # 显示结果
            self.show_result(data, used_ocr)
            self.save_btn.config(state='normal')

            # 扣配额
            allowed, msg = self.checker.check_and_increment(os.path.basename(pdf_path))
            self.log(msg)
        except Exception as e:
            self.log(f"✗ 识别失败: {e}")
            messagebox.showerror("错误", f"识别失败:\n{str(e)}")

    def show_result(self, data, used_ocr=False):
        """显示识别结果"""
        lines = []
        lines.append(f"{'='*50}")
        lines.append("发票识别结果" + ("（OCR识别）" if used_ocr else ""))
        lines.append(f"{'='*50}")
        lines.append(f"发票类型:   {data.get('invoice_type') or '未识别'}")
        lines.append(f"发票抬头:   {data.get('invoice_title') or '未识别'}")
        lines.append(f"发票代码:   {data.get('invoice_code') or '未识别'}")
        lines.append(f"发票号码:   {data.get('invoice_number') or '未识别'}")
        lines.append(f"开票日期:   {data.get('issue_date') or '未识别'}")
        lines.append("")
        lines.append(f"购买方名称: {data.get('buyer_name') or '未识别'}")
        lines.append(f"购买方税号: {data.get('buyer_taxid') or '未识别'}")
        lines.append(f"销售方名称: {data.get('seller_name') or '未识别'}")
        lines.append(f"销售方税号: {data.get('seller_taxid') or '未识别'}")
        lines.append("")
        lines.append(f"价税合计:   ¥{data.get('total_amount') or '未识别'}"
                     + (f"  ({data.get('total_amount_upper')})" if data.get('total_amount_upper') else ""))
        lines.append(f"校验码:     {data.get('check_code') or '未识别'}")
        if data.get('remark'):
            lines.append(f"备注:       {data['remark']}")
        if data.get('drawer'):
            lines.append(f"开票人:     {data['drawer']}")
        lines.append("")
        lines.append(f"项目明细 ({len(data.get('items', []))} 项):")
        if data.get('items'):
            for i, item in enumerate(data['items'], 1):
                lines.append(f"  {i}. {item['name']}"
                             + (f" {item['qty']}{item['unit']}" if item.get('qty') else "")
                             + (f" 金额:{item['amount']}" if item.get('amount') else "")
                             + (f" 税率:{item['tax_rate']}" if item.get('tax_rate') else "")
                             + (f" 税额:{item['tax']}" if item.get('tax') else ""))
        else:
            lines.append("  （未识别到明细行）")
        lines.append(f"{'='*50}")

        self.result_text.config(state='normal')
        self.result_text.delete('1.0', tk.END)
        self.result_text.insert('1.0', '\n'.join(lines))
        self.result_text.config(state='disabled')

    def copy_result(self):
        try:
            content = self.result_text.get('1.0', tk.END).strip()
            if content:
                self.win.clipboard_clear()
                self.win.clipboard_append(content)
                self.log("识别结果已复制到剪贴板")
                messagebox.showinfo("成功", "识别结果已复制到剪贴板")
        except Exception as e:
            messagebox.showerror("错误", str(e))

    def save_excel(self):
        """保存识别结果为Excel"""
        if not self.result_data:
            return
        output_dir = self.out_var.get().strip()
        if not output_dir or not os.path.exists(output_dir):
            try:
                os.makedirs(output_dir, exist_ok=True)
            except Exception:
                output_dir = os.path.expanduser("~")
                self.out_var.set(output_dir)

        number = self.result_data.get('invoice_number') or '未知'
        date = self.result_data.get('issue_date') or '无日期'
        output_path = os.path.join(output_dir, f"发票_{number}_{date}.xlsx")
        try:
            save_invoice_excel(self.result_data, output_path, self.log)
            messagebox.showinfo("成功", f"识别结果已保存:\n{output_path}")
        except Exception as e:
            self.log(f"保存Excel失败: {e}")
            messagebox.showerror("错误", f"保存失败:\n{str(e)}")
